from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
CT='http://schemas.openxmlformats.org/package/2006/content-types'
PR='http://schemas.openxmlformats.org/package/2006/relationships'
ET.register_namespace('w',W); ET.register_namespace('r',R)

# Các dòng mốc được giữ ổn định để generator.py thay số liệu theo từng tuần.
BODY = [
('title','VIỆN KIỂM SÁT NHÂN DÂN TỈNH NINH BÌNH'),
('title','CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM'),
('title','Độc lập - Tự do - Hạnh phúc'),
('title','BÁO CÁO'),
('title','Tình hình tội phạm và kết quả công tác tuần'),
('center','(Từ ngày 26/8/2026 đến ngày 09/9/2026)'),
('h1','I. TÌNH HÌNH TỘI PHẠM VÀ VI PHẠM'),
('h2','1. Tình hình tội phạm'),
('p','Trong tuần, Cơ quan điều tra đã khởi tố 37 vụ 100 bị can, gồm:'),
('p','1.1. Tội phạm xâm phạm an ninh quốc gia: Không.'),
('p','1.2. Tội phạm về trật tự xã hội:'),
('p','1.3. Tội phạm về kinh tế, sở hữu, môi trường:'),
('p','1.4. Tội phạm về ma tuý:'),
('p','1.5. Tội phạm về tham nhũng, chức vụ:'),
('p','1.6. Tội xâm phạm hoạt động tư pháp: Không.'),
('h2','2. Tình hình tai nạn, cháy nổ'),
('p','- Tai nạn giao thông:'),('p','- Tai nạn lao động, cháy nổ:'),
('h1','II. CÔNG TÁC CÔNG TỐ, KIỂM SÁT HOẠT ĐỘNG TƯ PHÁP'),
('h2','1. Công tác công tố, kiểm sát việc tiếp nhận, giải quyết tố giác, tin báo về tội phạm và kiến nghị khởi tố'),
('p','- Thực hành quyền công tố, kiểm sát việc tiếp nhận, giải quyết:'),
('p','- Chuyển giải quyết theo thẩm quyền:'),('p','- Đã giải quyết: 62 tin, gồm:'),
('p','+ Khởi tố vụ án hình sự:'),('p','+ Không khởi tố vụ án hình sự:'),('p','+ Tạm đình chỉ: 8 tin.'),
('p','- Đang giải quyết: 258 tin (trong hạn luật định).'),
('p','- Số tố giác, tin báo về tội phạm, kiến nghị khởi tố tạm đình chỉ giải quyết đến cuối kỳ báo cáo: 517 tin.'),
('h3','* Hoạt động thực hành quyền công tố, kiểm sát việc tiếp nhận, giải quyết tố giác, tin báo về tội phạm và kiến nghị khởi tố:'),
('p','- Ban hành 58 yêu cầu kiểm tra, xác minh giải quyết nguồn tin tội phạm;'),
('p','- Trực tiếp kiểm sát việc tiếp nhận, giải quyết tố giác, tin báo về tội phạm và kiến nghị khởi tố'),
('p','* Vi phạm trong hoạt động tư pháp: Một số hồ sơ tin báo'),
('p','* Những biện pháp tác động của Viện kiểm sát và việc chấp hành của cơ quan, cá nhân liên quan: Ban hành 2 Kiến nghị vi phạm'),
('p','* Kiến nghị cơ quan, tổ chức hữu quan áp dụng biện pháp phòng ngừa tội phạm và vi phạm pháp luật: Không.'),
('h2','2. Công tác công tố, kiểm sát việc khởi tố, điều tra và truy tố vụ án hình sự'),
('h3','2.1. Công tố, quyết định áp dụng, thay đổi, hủy bỏ các biện pháp ngăn chặn: Không.'),
('h3','2.2. Công tố, kiểm sát điều tra vụ án hình sự'),
('p','- Công tố, kiểm sát điều tra 819 vụ 1549 bị can'),('p','- Cơ quan điều tra đã giải quyết:'),
('p','+ Kết thúc điều tra đề nghị truy tố:'),('p','+ Đình chỉ: 3 vụ 4 bị can'),('p','+ Tạm đình chỉ: 5 vụ 3 bị can'),
('p','- Đang giải quyết: 704 vụ'),('p','- Tổng số án tạm đình chỉ điều tra tính đến cuối kỳ báo cáo: 755'),
('h3','* Hoạt động thực hành quyền công tố và kiểm sát điều tra:'),('p','- Ban hành 37 yêu cầu điều tra;'),('p','- Xác định án trọng điểm:'),
('p','* Vi phạm pháp luật điển hình trong hoạt động tư pháp giai đoạn điều tra:'),
('h3','* Những biện pháp tác động của Viện kiểm sát và việc chấp hành của cơ quan, cá nhân liên quan:'),
('p','- Ban hành 1 Kiến nghị vi phạm đối với Cơ quan CSĐT'),('p','- Ban hành 4 Kiến nghị phòng ngừa tội phạm'),
('h3','2.3. Công tố và kiểm sát hoạt động tư pháp trong giai đoạn truy tố'),
('p','- Công tố và kiểm sát việc giải quyết: 204'),('p','- Đã giải quyết: 132 vụ'),('p','+ Truy tố: 130 vụ'),
('p','+ Đình chỉ: 1 vụ 1 bị can'),('p','+ Nhập án: 1 vụ 1 bị can'),('p','- Đang giải quyết: 72 vụ'),
('p','- Tổng số án tạm đình chỉ điều tra tính đến cuối kỳ báo cáo: 1 vụ 1 bị can'),
('p','* Hoạt động công tố và kiểm sát hoạt động tư pháp trong giai đoạn truy tố: Không.'),('p','* Vi phạm trong hoạt động tư pháp: Không.'),
('h2','3. Công tác công tố và kiểm sát xét xử vụ án hình sự'),('h3','3.1. Công tố và kiểm sát xét xử sơ thẩm'),
('p','- Công tố và kiểm sát xét xử sơ thẩm 428'),('p','- Tòa án đã giải quyết: 130 vụ'),('p','- Đang giải quyết: 298 vụ'),
('p','- Tổng số án tạm đình chỉ điều tra tính đến cuối kỳ báo cáo: 0 vụ 0 bị cáo'),('p','- Phối hợp với Tòa án tổ chức 8 phiên tòa rút kinh nghiệm'),
('h3','3.2. Công tố và kiểm sát xét xử phúc thẩm'),('p','- Thụ lý 40 vụ 55 bị cáo'),('p','- Đã giải quyết: 9 vụ 22 bị cáo'),
('p','+ Xét xử: 5 vụ 7 bị cáo'),('p','+ Đình chỉ: 4 vụ 15 bị cáo'),('p','- Đang giải quyết: 31 vụ 33 bị cáo'),
('h3','3.3. Công tố và kiểm sát xét xử giám đốc thẩm, tái thẩm: Không.'),
('h2','4. Công tác kiểm sát việc tạm giữ, tạm giam, thi hành án hình sự'),('h3','4.1. Kiểm sát việc tạm giữ, tạm giam'),
('p','a) Kiểm sát việc tạm giữ:'),('p','Kiểm sát việc tạm giữ: 80 người (Số cũ 24 người; mới 56 người). Đã giải quyết: 52 người. Đang giải quyết: 28 người.'),
('p','b) Kiểm sát việc tạm giam:'),('p','- Kiểm sát việc tạm giam: 2.279 người (Số cũ 2.206 người; mới 73 người). Đã giải quyết: 216 người. Đang giải quyết: 2.063 người.'),
('h3','4.2. Kiểm sát thi hành án hình sự'),('p','- Thi hành án tử hình: Tổng số 36 người (số cũ). Ân giảm xuống chung thân: 01 người. Đang tạm giam, chưa có quyết định thi hành: 35 người.'),
('p','+ Tù có thời hạn: Tổng số thụ lý 242 người'),('p','+ Án treo, cải tạo không giam giữ: Tổng số thụ lý 3.771'),
('p','+ Cấm đảm nhiệm chức vụ, cấm hành nghề hoặc làm công việc nhất định 39'),('p','+ Tha tù trước thời hạn 3 người'),
('p','* Vi phạm trong hoạt động tư pháp: Cơ quan thi hành án hình sự Công an tỉnh'),
('p','* Những biện pháp tác động của Viện kiểm sát và việc chấp hành của cơ quan, cá nhân liên quan: Ban hành 1 Kiến nghị vi phạm đối với Cơ quan thi hành án hình sự'),
('h2','5. Công tác kiểm sát việc giải quyết vụ, việc dân sự, hôn nhân và gia đình'),
('p','- Thụ lý kiểm sát việc giải quyết 2.205 vụ, việc'),('p','- Tòa án đã giải quyết: 396 vụ, việc'),('p','+ Xét xử: 111 vụ'),
('p','+ Đình chỉ: 80 vụ, việc'),('p','+ Công nhận sự thỏa thuận của đương sự: 203'),('p','+ Mở phiên họp: 2 việc'),('p','- Đang giải quyết 1.809 vụ, việc'),
('p','- Thụ lý kiểm sát 2 vụ'),('p','- Thụ lý 3 vụ (Số mới)'),('p','- Ban hành 2 yêu cầu xác minh thu thập chứng cứ'),
('p','- Phối hợp với Tòa án tổ chức 6 phiên tòa rút kinh nghiệm'),('p','* Vi phạm trong hoạt động tư pháp: Tòa án vi phạm'),
('h2','6. Công tác kiểm sát việc giải quyết vụ việc kinh doanh, thương mại, lao động và những việc khác theo quy định của pháp luật'),
('p','- Thụ lý 107 vụ, việc'),('p','- Tòa án đã giải quyết: 14 vụ'),
('p','- Kiểm sát việc xem xét, quyết định áp dụng các biện pháp xử lý hành chính'),('p','- Kiểm sát việc giải quyết yêu cầu mở thủ tục phá sản:'),
('h2','7. Kiểm sát việc giải quyết các vụ án hành chính'),('p','- Thụ lý kiểm sát: 83 vụ'),('p','- Tòa án đã giải quyết: 10 vụ'),
('p','- Đang giải quyết: 73 vụ'),('p','- Thụ lý kiểm sát việc giải quyết: 6 vụ (Số cũ 5 vụ; số mới 1 vụ)'),
('p','- Ban hành 2 yêu cầu xác minh, thu thập chứng cứ'),('p','- Phối hợp với Tòa án tổ chức 3 phiên tòa rút kinh nghiệm'),
('h2','8. Công tác kiểm sát thi hành án dân sự, thi hành án hành chính'),('p','- Tổng số phải thi hành 8.638 việc'),
('p','- Đã kết thúc thi hành án: 462 việc'),('p','- Đang thi hành: 8.176 việc'),('p','- Trực tiếp kiểm sát 1 cuộc tại Phòng THADS'),
('p','* Vi phạm trong hoạt động tư pháp: Chậm xác minh'),('p','- Ban hành 4 Kiến nghị vi phạm đối với THADS tỉnh'),
('h2','9. Công tác giải quyết khiếu nại, tố cáo và kiểm sát việc giải quyết khiếu nại, tố cáo trong hoạt động tư pháp'),
('p','9.1. Công tác tiếp công dân:'),('p','- Viện kiểm sát đã tiếp nhận 66 đơn 64 việc'),('p','+ Đơn thuộc thẩm quyền giải quyết:'),
('p','+ Đơn thuộc trách nhiệm kiểm sát việc giải quyết:'),('p','+ Đơn không thuộc thẩm quyền giải quyết và kiểm sát việc giải quyết:'),
('p','- Viện kiểm sát phải xử lý: 15 đơn 13 việc'),('p','- Đã xử lý: 2 đơn 2 việc'),('p','- Đang xử lý: 13 đơn 11 việc'),
('p','- Viện kiểm sát phải xử lý: 1 đơn 1 việc (số mới)'),('p','- Đã xử lý: 1 đơn 1 việc (Chuyển Tòa án)'),
('h2','10. Công tác giải quyết đơn yêu cầu bồi thường thiệt hại trong hoạt động tố tụng hình sự: Không.'),
('h2','11. Công tác giải quyết đơn đề nghị Giám đốc thẩm, tái thẩm'),('h2','12. Công tác kiểm sát bản án, quyết định giám đốc thẩm, tái thẩm: Không.'),
('h1','III. CÔNG TÁC KHÁC'),
('h2','1. Công tác quản lý, chỉ đạo, điều hành'),
('p','- Tiếp tục quán triệt, triển khai các Chỉ thị, Kế hoạch công tác của VKSND tối cao và VKSND tỉnh Ninh Bình; nâng cao trách nhiệm người đứng đầu, kỷ luật, kỷ cương công vụ.'),
('h2','2. Công tác tổ chức cán bộ'),('p','- Tiếp tục thực hiện các nhiệm vụ về tổ chức bộ máy, vị trí việc làm, đào tạo, bồi dưỡng và công tác Đảng theo chương trình công tác.'),
('h2','3. Công tác tham mưu, tổng hợp, tuyên truyền, thi đua khen thưởng, báo cáo thống kê, tài chính, công nghệ thông tin'),
('p','- Tiếp tục thực hiện các nhiệm vụ tham mưu tổng hợp, báo cáo thống kê, tuyên truyền, thi đua, tài chính và chuyển đổi số theo kế hoạch.'),
('h2','4. Công tác thanh tra, kiểm tra, hướng dẫn nghiệp vụ, trả lời thỉnh thị'),('p','- Tiếp tục tăng cường thanh tra, kiểm tra, hướng dẫn nghiệp vụ và chấp hành kỷ luật, kỷ cương.'),
('h2','5. Công tác sơ kết, tổng kết, tập huấn, xây dựng pháp luật: Không.'),('h2','6. Công tác phối hợp và đối ngoại: Không.'),
('h1','IV. NHIỆM VỤ CÔNG TÁC TRỌNG TÂM TUẦN SAU'),
('p','- Tiếp tục thực hiện các Nghị quyết, Chỉ thị, Kế hoạch của VKSND tối cao và Kế hoạch công tác của VKSND tỉnh; tăng cường quản lý, thanh tra, kiểm tra, hướng dẫn và kịp thời giải quyết các vấn đề phát sinh.'),
('p','Trên đây là báo cáo kết quả công tác tuần và nhiệm vụ trọng tâm công tác tuần tới; Viện kiểm sát nhân dân tỉnh Ninh Bình báo cáo Viện kiểm sát nhân dân tối cao theo dõi, chỉ đạo./.'),
('p','Nơi nhận:\n- VKSND tối cao;\n- Văn phòng Tỉnh uỷ;\n- Ban Nội chính;\n- Lãnh đạo VKS tỉnh;\n- Các phòng thuộc VKS tỉnh;\n- Các VKS khu vực;\n- Lưu VT, VP.'),
('right','KT. VIỆN TRƯỞNG\nPHÓ VIỆN TRƯỞNG\nNguyễn Quốc Phương')
]

def _add_footnote_part(path):
    tmp=Path(str(path)+'.tmp')
    with ZipFile(path,'r') as zin, ZipFile(tmp,'w',ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data=zin.read(item.filename)
            if item.filename=='[Content_Types].xml':
                root=ET.fromstring(data)
                if not any(x.attrib.get('PartName')=='/word/footnotes.xml' for x in root):
                    x=ET.SubElement(root,'{%s}Override'%CT); x.set('PartName','/word/footnotes.xml'); x.set('ContentType','application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml')
                data=ET.tostring(root,encoding='utf-8',xml_declaration=True)
            elif item.filename=='word/_rels/document.xml.rels':
                root=ET.fromstring(data)
                if not any(x.attrib.get('Type','').endswith('/footnotes') for x in root):
                    ids=[int(x.attrib.get('Id','rId0').replace('rId','') or 0) for x in root if x.attrib.get('Id','').startswith('rId') and x.attrib.get('Id','')[3:].isdigit()]
                    x=ET.SubElement(root,'{%s}Relationship'%PR); x.set('Id','rId%d'%(max(ids or [0])+1)); x.set('Type','http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes'); x.set('Target','footnotes.xml')
                data=ET.tostring(root,encoding='utf-8',xml_declaration=True)
            zout.writestr(item,data)
        root=ET.Element('{%s}footnotes'%W)
        for fid,typ in [(-1,'separator'),(0,'continuationSeparator')]:
            fn=ET.SubElement(root,'{%s}footnote'%W); fn.set('{%s}id'%W,str(fid)); fn.set('{%s}type'%W,typ)
            p=ET.SubElement(fn,'{%s}p'%W); r=ET.SubElement(p,'{%s}r'%W); ET.SubElement(r,'{%s}%s'%(W,typ))
        for fid in range(1,40):
            fn=ET.SubElement(root,'{%s}footnote'%W); fn.set('{%s}id'%W,str(fid)); p=ET.SubElement(fn,'{%s}p'%W)
            r=ET.SubElement(p,'{%s}r'%W); rp=ET.SubElement(r,'{%s}rPr'%W); st=ET.SubElement(rp,'{%s}rStyle'%W); st.set('{%s}val'%W,'FootnoteReference'); ref=ET.SubElement(r,'{%s}footnoteRef'%W)
            r2=ET.SubElement(p,'{%s}r'%W); t=ET.SubElement(r2,'{%s}t'%W); t.text=' '
        zout.writestr('word/footnotes.xml',ET.tostring(root,encoding='utf-8',xml_declaration=True))
    tmp.replace(path)

def ensure_template(path):
    path=Path(path)
    if path.exists(): return path
    d=Document()
    sec=d.sections[0]; sec.top_margin=Cm(2); sec.bottom_margin=Cm(2); sec.left_margin=Cm(3); sec.right_margin=Cm(2)
    normal=d.styles['Normal']; normal.font.name='Times New Roman'; normal.font.size=Pt(14)
    for kind,text in BODY:
        p=d.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.15
        r=p.add_run(text); r.font.name='Times New Roman'; r.font.size=Pt(14)
        if kind in ('title','h1','h2','h3','right'): r.bold=True
        if kind in ('title','center'): p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        if kind=='right': p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        if kind in ('h1','h2','h3'): p.paragraph_format.keep_with_next=True
    d.save(path); _add_footnote_part(path); return path
