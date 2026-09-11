import io, os, tempfile, traceback, zipfile
from pathlib import Path
from flask import Flask, request, send_file, render_template_string, make_response
from werkzeug.exceptions import RequestEntityTooLarge, ClientDisconnected, BadRequest
from core_overlay import update_core
from office_overlay import overlay_office
from excel_overlay import apply_excel_activity_overlay
from footnote_fix import fix_footnote_format
from docx import Document

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 60 * 1024 * 1024
BASE = Path(__file__).resolve().parent
TEMPLATE = Path(os.environ.get('REPORT_TEMPLATE_PATH', str(BASE/'report_template.docx')))

HTML='''<!doctype html><html lang="vi"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Tạo báo cáo tuần đầy đủ</title><style>*{box-sizing:border-box}body{font-family:Arial,sans-serif;background:#f4f6f8;margin:0;color:#1f2937}.wrap{max-width:820px;margin:24px auto;padding:14px}.card{background:#fff;border-radius:14px;padding:24px;box-shadow:0 8px 28px #0001}.brand{font-weight:700;color:#991b1b;font-size:22px}.sub,.note{line-height:1.55}.sub{color:#4b5563;margin:8px 0 20px}.field{margin:16px 0}.field label{display:block;font-weight:700;margin-bottom:7px}.field input{width:100%;border:1px solid #d1d5db;padding:12px;border-radius:8px}.fn{font-size:13px;color:#166534;margin-top:6px}.btn{width:100%;border:0;border-radius:9px;padding:15px;background:#991b1b;color:#fff;font-weight:700;font-size:16px}.status{display:none;margin-top:12px;padding:10px;background:#eff6ff;border-radius:8px}.note{margin-top:16px;padding:12px;background:#ecfdf5;border-radius:8px;font-size:13px}.err{background:#fef2f2;color:#991b1b;padding:12px;border-radius:8px;margin-bottom:15px}</style></head><body><div class="wrap"><div class="card"><div class="brand">TẠO BÁO CÁO TUẦN — BẢN MỚI ĐẦY ĐỦ</div><div class="sub">Luôn giữ nguyên mẫu báo cáo tỉnh làm khung; Excel cập nhật toàn bộ nhóm nghiệp vụ, Phòng 7 và Phòng 9 bổ sung nội dung chuyên môn tương ứng.</div>{% if error %}<div class="err">{{error}}</div>{% endif %}<form id="f" method="post" enctype="multipart/form-data"><div class="field"><label>1. Bảng Excel báo cáo tuần</label><input id="excel" name="excel" type="file" accept=".xlsx" required><div class="fn" id="excelName"></div></div><div class="field"><label>2. Báo cáo Phòng 7</label><input id="p7" name="p7" type="file" accept=".doc,.docx" required><div class="fn" id="p7Name"></div></div><div class="field"><label>3. Báo cáo Phòng 9</label><input id="p9" name="p9" type="file" accept=".doc,.docx" required><div class="fn" id="p9Name"></div></div><button class="btn" id="b">TẠO VÀ TẢI BÁO CÁO WORD ĐẦY ĐỦ</button><div class="status" id="s">Đang tạo báo cáo đầy đủ từ 3 nguồn…</div></form><div class="note"><b>Nguyên tắc:</b> không tạo lại tài liệu trắng; hệ thống sao chép nguyên mẫu Word chuẩn trước, sau đó cập nhật số liệu và nội dung. Số liệu mới màu đỏ; footnote là footnote thật, ký hiệu số mũ; có kiểm tra cấu trúc tài liệu để phát hiện trường hợp đầu ra bị thiếu nghiêm trọng, không phụ thuộc tên/đánh số tiêu đề cố định.</div></div></div><script>['excel','p7','p9'].forEach(x=>document.getElementById(x).onchange=e=>document.getElementById(x+'Name').textContent=e.target.files[0]?'Đã chọn: '+e.target.files[0].name:'');document.getElementById('f').onsubmit=()=>{document.getElementById('b').disabled=true;document.getElementById('s').style.display='block'}</script></body></html>'''

def render_page(error=None,status=200):
    r=make_response(render_template_string(HTML,error=error),status)
    r.headers['Cache-Control']='no-store'
    return r

def valid_doc(name): return Path(name or '').suffix.lower() in {'.doc','.docx'}
def valid_xlsx(path):
    try:
        with zipfile.ZipFile(path) as z: return 'xl/workbook.xml' in z.namelist()
    except Exception: return False
def valid_docx(path):
    try:
        with zipfile.ZipFile(path) as z: return 'word/document.xml' in z.namelist() and 'word/styles.xml' in z.namelist()
    except Exception: return False

def structural_guard(path, template_path):
    """Chỉ chặn khi tài liệu thực sự bị mất cấu trúc đáng kể.
    So sánh tương đối với chính mẫu Word, không dùng ngưỡng tuyệt đối lớn hơn mẫu.
    """
    doc=Document(path)
    base=Document(template_path)
    para_count=len(doc.paragraphs)
    base_para_count=max(1,len(base.paragraphs))
    table_count=len(doc.tables)
    base_table_count=len(base.tables)
    text_len=sum(len(p.text or '') for p in doc.paragraphs)
    base_text_len=max(1,sum(len(p.text or '') for p in base.paragraphs))
    problems=[]

    min_para_count=max(1, int(base_para_count * 0.55))
    if para_count < min_para_count and (base_para_count - para_count) >= 5:
        problems.append(f'số đoạn văn giảm bất thường ({para_count}/{base_para_count})')

    if base_table_count:
        min_table_count=max(1, (base_table_count + 1) // 2)
        if table_count < min_table_count and (base_table_count - table_count) >= 2:
            problems.append(f'số bảng bị thiếu ({table_count}/{base_table_count})')

    min_text_len=max(300, int(base_text_len * 0.45))
    if text_len < min_text_len and (base_text_len - text_len) >= 500:
        problems.append(f'nội dung văn bản giảm bất thường ({text_len}/{base_text_len} ký tự)')
    return problems

@app.get('/health')
def health():
    return {'ok':valid_docx(TEMPLATE),'version':'complete-new-v3-relative-structural-guard','inputs':3,'template_first':True,'structural_guard':True,'red_excel_numbers':True,'footnote_superscript':True}

@app.errorhandler(RequestEntityTooLarge)
def too_large(e): return render_page('Tổng dung lượng 3 tệp vượt quá 60 MB.',413)

@app.route('/',methods=['GET','POST'])
def index():
    if request.method=='GET': return render_page()
    try:
        ex,p7,p9=request.files.get('excel'),request.files.get('p7'),request.files.get('p9')
        if not all(x and x.filename for x in (ex,p7,p9)):
            return render_page('Vui lòng chọn đủ 3 tệp: Excel, Phòng 7 và Phòng 9.',400)
        if Path(ex.filename).suffix.lower()!='.xlsx' or not valid_doc(p7.filename) or not valid_doc(p9.filename):
            return render_page('Định dạng không đúng: Excel phải .xlsx; Phòng 7 và Phòng 9 phải .doc hoặc .docx.',400)
        if not valid_docx(TEMPLATE): return render_page('Mẫu Word chuẩn chưa sẵn sàng.',500)
        with tempfile.TemporaryDirectory() as td0:
            td=Path(td0)
            ep=td/'input.xlsx'; p7p=td/('p7'+Path(p7.filename).suffix.lower()); p9p=td/('p9'+Path(p9.filename).suffix.lower())
            out=td/'Bao_cao_tuan_DAY_DU_VKSND_tinh_Ninh_Binh.docx'
            ex.save(ep); p7.save(p7p); p9.save(p9p)
            if not valid_xlsx(ep): return render_page('File Excel không hợp lệ.',400)
            update_core(str(TEMPLATE),str(ep),str(out))
            overlay_office(str(out),str(p7p),str(p9p))
            apply_excel_activity_overlay(str(out),str(ep))
            fix_footnote_format(str(out))
            if not valid_docx(out): raise RuntimeError('File Word đầu ra không hợp lệ.')
            problems=structural_guard(out,TEMPLATE)
            if problems: raise RuntimeError('Báo cáo đầu ra bị mất cấu trúc: '+ '; '.join(problems))
            data=out.read_bytes()
        return send_file(io.BytesIO(data),as_attachment=True,download_name='Bao_cao_tuan_DAY_DU_VKSND_tinh_Ninh_Binh.docx',mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    except (ClientDisconnected,BadRequest): return render_page('Kết nối tải tệp bị gián đoạn. Vui lòng thử lại.',400)
    except Exception as e:
        traceback.print_exc(); return render_page('Lỗi xử lý: '+str(e),500)

if __name__=='__main__': app.run(host='0.0.0.0',port=int(os.environ.get('PORT','10000')))
