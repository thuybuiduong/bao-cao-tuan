import re, shutil, zipfile
from copy import deepcopy
from pathlib import Path
import openpyxl
from docx import Document
from docx.shared import RGBColor
from lxml import etree

RED=RGBColor(255,0,0)
TARGETS={
'Tin báo':'hoạt động thực hành quyền công tố, kiểm sát việc tiếp nhận',
'KSĐT':'hoạt động thực hành quyền công tố và kiểm sát điều tra',
'Truy tố':'hoạt động công tố và kiểm sát hoạt động tư pháp trong giai đoạn truy tố',
'XXST':'hoạt động công tố và kiểm sát hoạt động tư pháp trong giai đoạn xét xử',
'THAHS':'hoạt động kiểm sát thi hành án hình sự',
'DS.HNGĐ':'hoạt động kiểm sát việc giải quyết vụ việc dân sự',
'KDTM':'hoạt động kiểm sát việc giải quyết vụ việc kinh doanh',
'QĐ AD BPXLHC tại Tòa':'hoạt động kiểm sát giải quyết vụ, việc kinh doanh',
'HC':'hoạt động kiểm sát việc giải quyết vụ án hành chính',
'THADS':'hoạt động kiểm sát thi hành án dân sự',
'Khiếu tố':'* hoạt động kiểm sát:'}
LABELS={'yc điều tra':'ban hành yêu cầu điều tra','yc khởi tố vụ án':'yêu cầu khởi tố vụ án','yc khởi tố bị can':'yêu cầu khởi tố bị can','yc thay đổi quyết định khởi tố bị can':'yêu cầu thay đổi quyết định khởi tố bị can','yc cqđt chỉ đạo đẩy nhanh các hoạt động điều tra vụ án':'yêu cầu CQĐT đẩy nhanh hoạt động điều tra','yc cqđt rà soát, giải quyết đối với 01 vụ án tạm đình chỉ đã hết thời hiệu truy cứu trách nhiệm hình sự':'yêu cầu CQĐT rà soát, giải quyết vụ án tạm đình chỉ đã hết thời hiệu','yêu cầu kiểm tra, xác minh giải quyết nguồn tin tội phạm':'ban hành yêu cầu kiểm tra, xác minh giải quyết nguồn tin tội phạm','yêu cầu giải quyết nguồn tin tội phạm':'yêu cầu giải quyết nguồn tin tội phạm','yêu cầu phục hồi giải quyết nguồn tin về tội phạm':'yêu cầu phục hồi giải quyết nguồn tin về tội phạm','án điểm/ tttg hỏi cung':'xác định án trọng điểm/trực tiếp, tham gia hỏi cung','phiên tòa rút kinh nghiệm':'phối hợp tổ chức phiên tòa rút kinh nghiệm','phiên tòa trình chiếu tài liệu, chứng cứ được số hóa':'phối hợp tổ chức phiên tòa trình chiếu tài liệu, chứng cứ được số hóa','lãnh đạo xx':'Lãnh đạo tham gia xét xử','số hóa hồ sơ':'số hóa hồ sơ','yêu cầu xác minh thu thập chứng cứ':'yêu cầu xác minh, thu thập chứng cứ','bc án sơ đồ tư duy':'báo cáo án bằng sơ đồ tư duy','thông báo rkn':'thông báo rút kinh nghiệm'}
MEAS={'vụ','bị can','bị cáo','việc','đơn','người','tiền','đồng','tin'}
def norm(s): return re.sub(r'\s+',' ',str(s or '')).strip().lower()
def num(v): return isinstance(v,(int,float)) and not isinstance(v,bool)
def fmt(v):
    if isinstance(v,float) and v.is_integer(): v=int(v)
    return f'{v:,}'.replace(',','.') if isinstance(v,(int,float)) else str(v)
def unit(raw):
    s=' '.join(str(raw or '').split())
    if not s or norm(s)=='tổng': return ''
    m=re.match(r'KV\s*(\d+)',s,re.I)
    if m: return 'KV'+m.group(1)
    m=re.match(r'Phòng\s*([\w]+)',s,re.I)
    return f'Phòng {m.group(1)}' if m else s.split('\n')[0]
def groups(ws):
    ss=[c for c in range(3,ws.max_column+1) if ws.cell(1,c).value not in (None,'')]; out=[]
    for i,s in enumerate(ss):
        e=ss[i+1]-1 if i+1<len(ss) else ws.max_column
        for rg in ws.merged_cells.ranges:
            if rg.min_row==1 and rg.max_row==1 and rg.min_col==s: e=max(e,rg.max_col)
        out.append((s,e,ws.cell(1,s).value))
    return out
def total_group(ws): return next((g for g in groups(ws) if norm(g[2])=='tổng'),None)
def hrow(ws,g,limit=30):
    gs,ge,_=g; best=(0,None)
    for r in range(2,min(limit,8,ws.max_row)+1):
        sc=sum(1 for c in range(gs,ge+1) if norm(ws.cell(r,c).value) in MEAS)
        if sc>best[0]: best=(sc,r)
    return best[1]
def measure(ws,r,c):
    v=ws.cell(r,c).value if r else None
    return str(v).strip() if isinstance(v,str) and v.strip() else None
def rowno(ws,label):
    want=norm(label)
    for r in range(1,min(ws.max_row,120)+1):
        got=norm(ws.cell(r,1).value)
        if got==want or (want and got.startswith(want)): return r
    return None
def vals(ws,label):
    tg=total_group(ws); r=rowno(ws,label)
    if not tg or not r: return []
    hr=hrow(ws,tg,r); return [(measure(ws,hr,c),ws.cell(r,c).value) for c in range(tg[0],tg[1]+1) if num(ws.cell(r,c).value)]
def pair(ws,label):
    vu=bc=0
    for m,v in vals(ws,label):
        nm=norm(m)
        if 'bị can' in nm or 'bị cáo' in nm: bc=v
        elif 'vụ' in nm or not nm: vu=v if vu==0 else vu
    return vu,bc
def act_start(ws):
    for r in range(1,ws.max_row+1):
        if 'hoạt động kiểm sát' in norm(' '.join(str(ws.cell(r,c).value or '') for c in range(1,4))): return r
    return None
def act_rows(ws):
    st=act_start(ws); tg=total_group(ws)
    if not st or not tg: return []
    ts,te,_=tg; hr=hrow(ws,tg,st); tcols=[(c,measure(ws,hr,c)) for c in range(ts,te+1) if measure(ws,hr,c) or te==ts] or [(ts,None)]
    cgs=[g for g in groups(ws) if g[0]<ts]; out=[]; blanks=0
    for r in range(st+1,min(ws.max_row,st+100)+1):
        lab=ws.cell(r,1).value or ws.cell(r,2).value
        if lab in (None,''):
            blanks+=1
            if blanks>=3: break
            continue
        blanks=0; lab=str(lab).strip()
        if norm(lab).startswith('vi phạm'): continue
        contrib=[]
        for gs,ge,gn in cgs:
            un=unit(gn); gh=hrow(ws,(gs,ge,gn),st) or hr
            if not un: continue
            for c in range(gs,ge+1):
                v=ws.cell(r,c).value
                if num(v) and v!=0: contrib.append((un,measure(ws,gh,c),v))
        ms=[]
        for c,m in tcols:
            v=ws.cell(r,c).value
            if num(v) and v!=0: ms.append((m,v))
        if not ms and contrib:
            d={}
            for _,m,v in contrib: d[m]=d.get(m,0)+v
            ms=list(d.items())
        if ms or contrib: out.append((lab,ms,contrib))
    return out
def heading(doc,needle):
    nn=norm(needle).replace('* ','')
    for i,p in enumerate(doc.paragraphs):
        if nn in norm(p.text).replace('* ',''): return i
    return None
def sect_end(doc,i):
    for j in range(i+1,len(doc.paragraphs)):
        if re.match(r'^\d+(?:\.\d+){0,2}\.\s',norm(doc.paragraphs[j].text)): return j
    return len(doc.paragraphs)
def hasfn(p): return bool(p._p.xpath('.//w:footnoteReference'))
def tokens(s):
    s=norm(LABELS.get(norm(s),s)); s=s.replace('tttg','trực tiếp tham gia').replace('rkn','rút kinh nghiệm').replace('xx','xét xử')
    stop={'ban','hành','công','tác','hoạt','động','kiểm','sát','việc','về','và','của','đối','với','theo','các'}
    return {x for x in re.findall(r'[a-zà-ỹđ]+',s) if len(x)>1 and x not in stop}
def find_para(paras,lab,meas=None):
    need=tokens(lab)
    if norm(lab)=='án điểm/ tttg hỏi cung': need={'án','trọng','điểm'} if 'vụ' in norm(meas) else {'hỏi','cung'}
    best=None
    for p in paras:
        have=set(re.findall(r'[a-zà-ỹđ]+',norm(p.text))); inter=len(need&have); cov=inter/max(1,len(need))
        if cov>=(1 if len(need)<=2 else .58):
            sc=(cov,inter)
            if best is None or sc>best[0]: best=(sc,p)
    return best[1] if best else None
def ftxt(contrib,meas):
    rows=[]; seen=set()
    for u,m,v in contrib:
        if meas is not None and norm(m)!=norm(meas): continue
        k=(u,v)
        if k not in seen: rows.append((u,v)); seen.add(k)
    if not rows and meas is not None:
        for u,m,v in contrib:
            k=(u,v)
            if k not in seen: rows.append((u,v)); seen.add(k)
    return '; '.join(f'{u} - {fmt(v)}' for u,v in rows)+'.' if rows else ''
def parts(lab,meas,v):
    nl=norm(lab); m=norm(meas); n=fmt(v)
    if nl=='án điểm/ tttg hỏi cung': return ('- Trực tiếp và tham gia hỏi cung ',n,' bị can.') if 'bị can' in m else ('- Xác định án trọng điểm: ',n,' vụ.')
    u=' bị can' if 'bị can' in m else ' bị cáo' if 'bị cáo' in m else ' vụ' if 'vụ' in m else ' việc' if 'việc' in m else ' đơn' if 'đơn' in m else ' người' if 'người' in m else ' đồng' if ('tiền' in m or 'đồng' in m) else (' lượt' if any(x in nl for x in ['trực tiếp kiểm sát','phiên tòa','số hóa']) else ' yêu cầu' if ('yêu cầu' in nl or nl.startswith('yc ')) else '')
    lab=LABELS.get(norm(lab),lab); return (f'- {lab[:1].upper()+lab[1:]}: ',n,u+'.')
def copy_ppr(src,dst):
    if src is not None and src._p.pPr is not None:
        cl=deepcopy(src._p.pPr)
        if dst._p.pPr is not None: dst._p.remove(dst._p.pPr)
        dst._p.insert(0,cl)
def insert_before(ref,prt,marker,src=None):
    p=ref.insert_paragraph_before(); copy_ppr(src,p); a,n,b=prt; p.add_run(a); r=p.add_run(n); r.font.color.rgb=RED; p.add_run(b)
    if marker: p.add_run(marker)
    return p
def red_text(p,text):
    for r in list(p._p.r_lst): p._p.remove(r)
    for x in re.split(r'(\d[\d\.]*)',text):
        if x=='': continue
        r=p.add_run(x)
        if re.fullmatch(r'\d[\d\.]*',x): r.font.color.rgb=RED
def patch_core(doc,wb):
    ch=0
    if 'KSĐT' in wb.sheetnames:
        ws=wb['KSĐT']; gq=pair(ws,'Giải quyết'); dt=pair(ws,'Đề nghị truy tố'); dc=pair(ws,'Đình chỉ'); td=pair(ws,'Tạm đình chỉ'); ton=pair(ws,'Tồn'); ck=pair(ws,'TĐC đến cuối kỳ')
        i=heading(doc,'Công tố, kiểm sát điều tra vụ án hình sự')
        if i is not None:
            for p in doc.paragraphs[i+1:sect_end(doc,i)]:
                nt=norm(p.text); new=None
                if 'cơ quan điều tra đã giải quyết' in nt: new=f'- Cơ quan điều tra đã giải quyết: {fmt(gq[0])} vụ {fmt(gq[1])} bị can, trong đó:'
                elif 'kết thúc điều tra đề nghị truy tố' in nt: new=f'+ Kết thúc điều tra đề nghị truy tố: {fmt(dt[0])} vụ {fmt(dt[1])} bị can;'
                elif 'tổng số án tạm đình chỉ điều tra tính đến cuối kỳ báo cáo' in nt: new=f'- Tổng số án tạm đình chỉ điều tra tính đến cuối kỳ báo cáo: {fmt(ck[0])} vụ {fmt(ck[1])} bị can.'
                elif nt.startswith('+ đình chỉ') and 'tạm đình chỉ' not in nt: new=f'+ Đình chỉ: {fmt(dc[0])} vụ {fmt(dc[1])} bị can;'
                elif nt.startswith('+ tạm đình chỉ'): new=f'+ Tạm đình chỉ: {fmt(td[0])} vụ {fmt(td[1])} bị can.'
                elif nt.startswith('- đang giải quyết') or nt.startswith('- còn đang') or nt.startswith('- tồn'): new=f'- Đang giải quyết: {fmt(ton[0])} vụ {fmt(ton[1])} bị can.'
                if new: red_text(p,new); ch+=1
    return ch
def add_footnotes(path,mapping):
    if not mapping: return
    src=Path(path); tmp=src.with_suffix('.finalfn.tmp.docx'); WNS='http://schemas.openxmlformats.org/wordprocessingml/2006/main'; W='{'+WNS+'}'
    with zipfile.ZipFile(src) as zin:
        doc=etree.fromstring(zin.read('word/document.xml')); fns=etree.fromstring(zin.read('word/footnotes.xml'))
        ids=[]
        for f in fns.findall(W+'footnote'):
            try: ids.append(int(f.get(W+'id')))
            except: pass
        nid=max([x for x in ids if x>=0] or [0])+1
        for mark,text in mapping.items():
            t=next((x for x in doc.findall('.//'+W+'t') if mark in (x.text or '')),None)
            if t is None: continue
            run=t.getparent(); t.text=(t.text or '').replace(mark,''); rpr=run.find(W+'rPr')
            if rpr is None: rpr=etree.Element(W+'rPr'); run.insert(0,rpr)
            rs=etree.SubElement(rpr,W+'rStyle'); rs.set(W+'val','FootnoteReference'); va=etree.SubElement(rpr,W+'vertAlign'); va.set(W+'val','superscript')
            fr=etree.SubElement(run,W+'footnoteReference'); fr.set(W+'id',str(nid))
            fn=etree.SubElement(fns,W+'footnote'); fn.set(W+'id',str(nid)); p=etree.SubElement(fn,W+'p'); ppr=etree.SubElement(p,W+'pPr'); ps=etree.SubElement(ppr,W+'pStyle'); ps.set(W+'val','FootnoteText')
            r1=etree.SubElement(p,W+'r'); r1p=etree.SubElement(r1,W+'rPr'); r1s=etree.SubElement(r1p,W+'rStyle'); r1s.set(W+'val','FootnoteReference'); etree.SubElement(r1,W+'footnoteRef')
            r2=etree.SubElement(p,W+'r'); tt=etree.SubElement(r2,W+'t'); tt.set('{http://www.w3.org/XML/1998/namespace}space','preserve'); tt.text=' '+text; nid+=1
        with zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as zout:
            for it in zin.infolist():
                data=etree.tostring(doc,xml_declaration=True,encoding='UTF-8',standalone='yes') if it.filename=='word/document.xml' else etree.tostring(fns,xml_declaration=True,encoding='UTF-8',standalone='yes') if it.filename=='word/footnotes.xml' else zin.read(it.filename)
                zout.writestr(it,data)
    shutil.move(tmp,src)
def apply_final_corrections(docx_path,xlsx_path):
    wb=openpyxl.load_workbook(xlsx_path,data_only=True); doc=Document(docx_path); fn={}; c=1; added=0; fadded=0; core=patch_core(doc,wb)
    for s,needle in TARGETS.items():
        if s not in wb.sheetnames: continue
        rows=act_rows(wb[s]); i=heading(doc,needle)
        if not rows or i is None: continue
        e=sect_end(doc,i); scope=list(doc.paragraphs[i+1:e]); ins=e
        for j in range(i+1,e):
            if norm(doc.paragraphs[j].text).startswith(('* vi phạm','* những biện pháp','* biện pháp')): ins=j; break
        ref=doc.paragraphs[ins] if ins<len(doc.paragraphs) else doc.paragraphs[-1]; src=next((p for p in scope if norm(p.text).startswith('-')),doc.paragraphs[i])
        for lab,measures,contrib in rows:
            for meas,v in measures:
                text=ftxt(contrib,meas); p=find_para(scope,lab,meas)
                if p is not None:
                    if text and not hasfn(p): mark=f'[[FFN{c}]]'; p.add_run(mark); fn[mark]=text; c+=1; fadded+=1
                else:
                    mark=f'[[FFN{c}]]' if text else None; np=insert_before(ref,parts(lab,meas,v),mark,src); scope.append(np); added+=1
                    if text: fn[mark]=text; c+=1; fadded+=1
    doc.save(docx_path); add_footnotes(docx_path,fn)
    return {'final_activity_added':added,'final_footnotes_added':fadded,'final_core_corrected':core}
