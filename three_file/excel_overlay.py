import re, shutil, zipfile
from copy import deepcopy
from pathlib import Path
import openpyxl
from docx import Document
from docx.shared import RGBColor
from lxml import etree

RED=RGBColor(255,0,0)
TARGETS={
'Tin báo':'hoạt động thực hành quyền công tố, kiểm sát việc tiếp nhận',
'KSĐT':'hoạt động thực hành quyền công tố và kiểm sát điều tra',
'Truy tố':'hoạt động công tố và kiểm sát hoạt động tư pháp trong giai đoạn truy tố',
'XXST':'hoạt động công tố và kiểm sát hoạt động tư pháp trong giai đoạn xét xử',
'THAHS':'hoạt động kiểm sát thi hành án hình sự',
'DS.HNGĐ':'hoạt động kiểm sát việc giải quyết vụ việc dân sự',
'KDTM':'hoạt động kiểm sát việc giải quyết vụ việc kinh doanh',
'QĐ AD BPXLHC tại Tòa':'hoạt động kiểm sát giải quyết vụ, việc kinh doanh',
'HC':'hoạt động kiểm sát việc giải quyết vụ án hành chính',
'THADS':'hoạt động kiểm sát thi hành án dân sự',
'Khiếu tố':'* hoạt động kiểm sát:'}
LABELS={
'yc điều tra':'yêu cầu điều tra','yc khởi tố vụ án':'yêu cầu khởi tố vụ án','yc khởi tố bị can':'yêu cầu khởi tố bị can',
'yc không khởi tố vụ án':'yêu cầu không khởi tố vụ án','yc không khởi tố bị can':'yêu cầu không khởi tố bị can',
'yc thay đổi quyết định khởi tố vụ án':'yêu cầu thay đổi quyết định khởi tố vụ án','yc thay đổi quyết định khởi tố bị can':'yêu cầu thay đổi quyết định khởi tố bị can',
'yc phong toả tài khoản':'yêu cầu phong tỏa tài khoản','yc phục hồi án tđc':'yêu cầu phục hồi vụ án tạm đình chỉ',
'yc giải quyết vụ án tđc':'yêu cầu rà soát, giải quyết vụ án tạm đình chỉ','yc truy nã':'yêu cầu truy nã',
'yc cqđt chỉ đạo đẩy nhanh các hoạt động điều tra vụ án':'yêu cầu cqđt đẩy nhanh hoạt động điều tra',
'yc cqđt rà soát, giải quyết đối với 01 vụ án tạm đình chỉ đã hết thời hiệu truy cứu trách nhiệm hình sự':'yêu cầu rà soát giải quyết vụ án tạm đình chỉ',
'yc bổ sung qđ ktva':'yêu cầu bổ sung quyết định khởi tố vụ án','yc bổ sung qđ ktbc':'yêu cầu bổ sung quyết định khởi tố bị can',
'án điểm/ tttg hỏi cung':'án điểm/ trực tiếp, tham gia hỏi cung','thông báo rkn':'thông báo rút kinh nghiệm','bc án sơ đồ tư duy':'báo cáo án bằng sơ đồ tư duy'}
MEASURE_WORDS={'vụ','bị can','bị cáo','việc','đơn','người','tiền','đồng','chung','tin','phiên tòa','phiên toà'}

def norm(s): return re.sub(r'\s+',' ',str(s or '')).strip().lower()
def is_num(v): return isinstance(v,(int,float)) and not isinstance(v,bool)
def fmt(v):
    if isinstance(v,float) and v.is_integer(): v=int(v)
    return f'{v:,}'.replace(',','.') if isinstance(v,(int,float)) else str(v)
def unit_name(raw):
    s=' '.join(str(raw or '').split())
    if not s or norm(s)=='tổng': return ''
    m=re.match(r'KV\s*(\d+)',s,re.I)
    if m: return f'VKS khu vực {m.group(1)}'
    m=re.match(r'Phòng\s*([\w]+)',s,re.I)
    if m: return f'Phòng {m.group(1)}'
    return s.split('\n')[0]
def find_activity_start(ws):
    for r in range(1,ws.max_row+1):
        txt=' '.join(str(ws.cell(r,c).value or '') for c in range(1,min(ws.max_column,4)+1))
        if 'hoạt động kiểm sát' in norm(txt): return r
    return None
def group_ranges(ws):
    starts=[]
    for c in range(3,ws.max_column+1):
        if ws.cell(1,c).value not in (None,''): starts.append(c)
    groups=[]
    for i,s in enumerate(starts):
        e=(starts[i+1]-1 if i+1<len(starts) else ws.max_column)
        for rg in ws.merged_cells.ranges:
            if rg.min_row==1 and rg.max_row==1 and rg.min_col==s:
                e=max(e,rg.max_col); break
        groups.append((s,e,ws.cell(1,s).value))
    return groups
def header_row(ws,total_group,start):
    ts,te,_=total_group; best=(0,None)
    for r in range(2,min(start or 8,8)+1):
        vals=[]
        for c in range(ts,te+1):
            v=ws.cell(r,c).value
            if isinstance(v,str) and norm(v) in MEASURE_WORDS: vals.append(v)
        if len(vals)>best[0]: best=(len(vals),r)
    return best[1]
def measure_at(ws,r,c):
    if not r: return None
    v=ws.cell(r,c).value
    return str(v).strip() if isinstance(v,str) and v.strip() else None

def activity_rows(ws):
    st=find_activity_start(ws)
    if not st: return []
    groups=group_ranges(ws); tg=next((g for g in groups if norm(g[2])=='tổng'),None)
    if not tg: return []
    ts,te,_=tg; hr=header_row(ws,tg,st)
    total_cols=[]
    for c in range(ts,te+1):
        m=measure_at(ws,hr,c)
        if m or te==ts: total_cols.append((c,m))
    if not total_cols: total_cols=[(ts,None)]
    contributor_groups=[g for g in groups if g[0]<ts]
    out=[]; blanks=0
    for r in range(st+1,min(ws.max_row,st+100)+1):
        label=ws.cell(r,1).value or ws.cell(r,2).value
        if label in (None,''):
            blanks+=1
            if blanks>=3: break
            continue
        blanks=0; label=str(label).strip(); nl=norm(label)
        if nl.startswith('vi phạm'): continue
        contrib=[]
        for gs,ge,gname in contributor_groups:
            u=unit_name(gname)
            if not u: continue
            for c in range(gs,ge+1):
                v=ws.cell(r,c).value
                if is_num(v) and v!=0:
                    m=measure_at(ws,hr,c)
                    contrib.append((u,m,v))
        measures=[]
        for c,m in total_cols:
            v=ws.cell(r,c).value
            if is_num(v) and v!=0: measures.append((m,v,False))
        if not measures and contrib:
            sums={}
            for _,m,v in contrib: sums[m]=sums.get(m,0)+v
            measures=[(m,v,True) for m,v in sums.items() if v]
        if measures or contrib: out.append((label,measures,contrib))
    return out

def human_label(label): return LABELS.get(norm(label),label)
def metric_parts(label,measure,value):
    nl=norm(label); n=fmt(value); m=norm(measure)
    if nl=='án điểm/ tttg hỏi cung':
        if 'bị can' in m: return ('- Trực tiếp và tham gia hỏi cung ',n,' bị can.')
        return ('- Xác định án trọng điểm: ',n,' vụ.')
    unit=''
    if 'bị can' in m: unit=' bị can'
    elif 'bị cáo' in m: unit=' bị cáo'
    elif 'vụ' in m: unit=' vụ'
    elif 'việc' in m: unit=' việc'
    elif 'đơn' in m: unit=' đơn'
    elif 'người' in m: unit=' người'
    elif 'tiền' in m or 'đồng' in m: unit=' đồng'
    lab=human_label(label)
    return (f'- {lab[:1].upper()+lab[1:]}: ',n,unit+'.')
def footnote_text(contrib,measure,value,derived=False):
    vals=[]; seen=set(); total=0
    for u,m,v in contrib:
        if measure is not None and norm(m)!=norm(measure): continue
        key=(u,m,v)
        if key in seen: continue
        seen.add(key); vals.append((u,v)); total+=v
    if not vals: return ''
    if not derived and is_num(value) and abs(total-value)>1e-9: return ''
    return '; '.join(f'{u} - {fmt(v)}' for u,v in vals)+'.'
def find_heading(doc,needle):
    nn=norm(needle).replace('* ','')
    for i,p in enumerate(doc.paragraphs):
        if nn in norm(p.text).replace('* ',''): return i
    return None
def section_end(doc,idx):
    for j in range(idx+1,len(doc.paragraphs)):
        t=norm(doc.paragraphs[j].text)
        if re.match(r'^\d+(?:\.\d+){0,2}\.\s',t): return j
    return len(doc.paragraphs)
def canon_text(s):
    t=norm(s)
    repl={
        ' rkn ':' rút kinh nghiệm ',' tđc ':' tạm đình chỉ ',' kl ':' kết luận ',' ts ':' tài sản ',
        ' qđ ':' quyết định ',' ktva ':' khởi tố vụ án ',' ktbc ':' khởi tố bị can ',
        ' tttg ':' trực tiếp tham gia ',' xx ':' xét xử '
    }
    t=' '+t+' '
    for a,b in repl.items(): t=t.replace(a,b)
    return ' '.join(t.split())

def match_tokens(label,measure=None):
    raw=norm(label)
    if raw=='án điểm/ tttg hỏi cung':
        return {'án','trọng','điểm'} if 'vụ' in norm(measure) else {'hỏi','cung'}
    h=canon_text(human_label(label))
    toks=set(re.findall(r'[a-zà-ỹđ]+',h))
    stop={'về','việc','công','tác','các','đối','với','trong','hoạt','động','của','và','theo','được','01','một'}
    return {x for x in toks if x not in stop and len(x)>1}

def matching_para(paras,label,measure=None):
    need=match_tokens(label,measure)
    if not need: return None
    best=None
    for p in paras:
        text=canon_text(p.text)
        have=set(re.findall(r'[a-zà-ỹđ]+',text))
        inter=len(need & have); cov=inter/max(1,len(need))
        threshold=1.0 if len(need)<=2 else 0.62
        if cov>=threshold:
            score=(cov,inter)
            if best is None or score>best[0]: best=(score,p)
    return best[1] if best else None

def insert_point(doc,idx,section_end_idx):
    for j in range(idx+1,section_end_idx):
        t=norm(doc.paragraphs[j].text)
        if (t.startswith('* vi phạm') or t.startswith('* những biện pháp') or
            t.startswith('* biện pháp') or t.startswith('* kiến nghị cơ quan')):
            return j
    return section_end_idx

def has_fn(p):
    return bool(p._p.xpath('.//w:footnoteReference'))
def copy_ppr(src,dst):
    if src is not None and src._p.pPr is not None:
        cl=deepcopy(src._p.pPr)
        if dst._p.pPr is not None: dst._p.remove(dst._p.pPr)
        dst._p.insert(0,cl)
def insert_before(ref,parts,marker=None,src=None):
    p=ref.insert_paragraph_before(); copy_ppr(src,p)
    a,n,b=parts; p.add_run(a); rr=p.add_run(n); rr.font.color.rgb=RED; p.add_run(b)
    if marker: p.add_run(marker)
    return p

def add_footnotes(path,mapping):
    if not mapping:return
    src=Path(path); tmp=src.with_suffix('.fn.tmp.docx'); WNS='http://schemas.openxmlformats.org/wordprocessingml/2006/main'; W='{'+WNS+'}'
    with zipfile.ZipFile(src,'r') as zin:
        doc=etree.fromstring(zin.read('word/document.xml')); fns=etree.fromstring(zin.read('word/footnotes.xml'))
        ids=[]
        for f in fns.findall(W+'footnote'):
            try: ids.append(int(f.get(W+'id')))
            except:pass
        nid=max([x for x in ids if x>=0] or [0])+1
        for marker,text in mapping.items():
            target=next((t for t in doc.findall('.//'+W+'t') if marker in (t.text or '')),None)
            if target is None: continue
            run=target.getparent(); target.text=(target.text or '').replace(marker,'')
            rpr=run.find(W+'rPr')
            if rpr is None:rpr=etree.Element(W+'rPr');run.insert(0,rpr)
            rs=etree.SubElement(rpr,W+'rStyle');rs.set(W+'val','FootnoteReference');va=etree.SubElement(rpr,W+'vertAlign');va.set(W+'val','superscript')
            fr=etree.SubElement(run,W+'footnoteReference');fr.set(W+'id',str(nid))
            fn=etree.SubElement(fns,W+'footnote');fn.set(W+'id',str(nid));p=etree.SubElement(fn,W+'p');ppr=etree.SubElement(p,W+'pPr');ps=etree.SubElement(ppr,W+'pStyle');ps.set(W+'val','FootnoteText')
            r1=etree.SubElement(p,W+'r');r1p=etree.SubElement(r1,W+'rPr');r1s=etree.SubElement(r1p,W+'rStyle');r1s.set(W+'val','FootnoteReference');etree.SubElement(r1,W+'footnoteRef')
            r2=etree.SubElement(p,W+'r');tt=etree.SubElement(r2,W+'t');tt.set('{http://www.w3.org/XML/1998/namespace}space','preserve');tt.text=' '+text;nid+=1
        with zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data=etree.tostring(doc,xml_declaration=True,encoding='UTF-8',standalone='yes') if item.filename=='word/document.xml' else etree.tostring(fns,xml_declaration=True,encoding='UTF-8',standalone='yes') if item.filename=='word/footnotes.xml' else zin.read(item.filename)
                zout.writestr(item,data)
    shutil.move(tmp,src)

def apply_excel_activity_overlay(docx_path,xlsx_path):
    wb=openpyxl.load_workbook(xlsx_path,data_only=True); doc=Document(docx_path); fnmap={}; counter=1; added=0; foot_added=0
    for sname,needle in TARGETS.items():
        if sname not in wb.sheetnames: continue
        rows=activity_rows(wb[sname]); idx=find_heading(doc,needle)
        if not rows or idx is None: continue
        end=section_end(doc,idx); scope=list(doc.paragraphs[idx+1:end])
        ins=insert_point(doc,idx,end); ref=doc.paragraphs[ins] if ins<len(doc.paragraphs) else doc.paragraphs[-1]
        fmt_src=next((p for p in scope if norm(p.text).startswith('-')),doc.paragraphs[idx])
        for label,measures,contrib in rows:
            for measure,value,derived in measures:
                existing=matching_para(scope,label,measure)
                ftxt=footnote_text(contrib,measure,value,derived)
                if existing is not None:
                    if ftxt and not has_fn(existing):
                        marker=f'[[EXFN{counter}]]'; existing.add_run(marker); fnmap[marker]=ftxt; foot_added+=1; counter+=1
                    continue
                marker=f'[[EXFN{counter}]]' if ftxt else None
                p=insert_before(ref,metric_parts(label,measure,value),marker,fmt_src); scope.append(p); added+=1
                if ftxt: fnmap[marker]=ftxt;foot_added+=1;counter+=1
    doc.save(docx_path); add_footnotes(docx_path,fnmap)
    return {'activity_metrics_added':added,'activity_footnotes_added':foot_added,'sheets_scanned':len(wb.sheetnames)}
